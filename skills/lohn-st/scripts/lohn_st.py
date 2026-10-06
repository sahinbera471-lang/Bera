#!/usr/bin/env python3
"""Lohn ST – Steuerberater-Ansicht der Monatsübersicht.

Blendet in der Excel-Monatsübersicht den oberen Teil (Übersicht über die
einzelnen Monate/Tage) aus, sodass nur noch der untere Teil für den
Steuerberater sichtbar ist (Urlaub/Konten, Monatsabrechnung, Weitere Angaben).

Es wird NICHTS gelöscht oder umgerechnet: Werte, Formeln, Formatierungen,
Auswahllisten und Makros bleiben unverändert. Das Skript setzt nur das
Attribut "ausgeblendet" an den Zeilen und den Druckbereich auf den sichtbaren
Teil – direkt im XML der Datei, damit keine Excel-Funktionen verloren gehen.
Die Originaldatei wird nicht angefasst; es entsteht eine Kopie.
"""

import argparse
import posixpath
import re
import sys
import zipfile
from xml.sax.saxutils import escape

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl fehlt – bitte 'pip install openpyxl' ausführen.")

# Überschriften der Abschnitte, ab denen alles sichtbar bleibt
MARKER_KONTEN = re.compile(r"^\s*(urlaub\s*/\s*konten|zeitkonto|arbeitszeitkonto)", re.I)
MARKER_ABRECHNUNG = re.compile(r"^\s*monatsabrechnung", re.I)


def finde_bereich(ws, mit_konten):
    """(erste sichtbare Zeile, letzte Zeile, letzte Spalte) des unteren Teils."""
    start = None
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and (
                    MARKER_ABRECHNUNG.search(c.value) or (mit_konten and MARKER_KONTEN.search(c.value))):
                start = c.row
                break
        if start:
            break
    if not start:
        return None
    letzte_zeile, letzte_spalte = start, 1
    for row in ws.iter_rows(min_row=start):
        for c in row:
            if c.value not in (None, ""):
                letzte_zeile = max(letzte_zeile, c.row)
                letzte_spalte = max(letzte_spalte, c.column)
    for rng in ws.merged_cells.ranges:  # verbundene Zellen (z. B. Bemerkungsfeld)
        if rng.min_row >= start:
            letzte_zeile = max(letzte_zeile, rng.max_row)
            letzte_spalte = max(letzte_spalte, rng.max_col)
    return start, letzte_zeile, letzte_spalte


# --------------------------------------------------------------------------
# XML-Bearbeitung
# --------------------------------------------------------------------------

def unescape(s):
    return (s.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
            .replace("&apos;", "'").replace("&amp;", "&"))


def blatt_pfade(z):
    """{Blattname: (Index, Pfad im ZIP)} aus workbook.xml und den Relationen."""
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    ziele = {}
    for m in re.finditer(r"<Relationship\b[^>]*>", rels):
        tag = m.group(0)
        rid = re.search(r'\bId="([^"]+)"', tag).group(1)
        ziel = re.search(r'\bTarget="([^"]+)"', tag).group(1)
        ziele[rid] = ziel.lstrip("/") if ziel.startswith("/") else posixpath.normpath("xl/" + ziel)
    blaetter = {}
    for i, m in enumerate(re.finditer(r"<(?:\w+:)?sheet\b[^>]*/>", wb_xml)):
        tag = m.group(0)
        name = re.search(r'\bname="([^"]*)"', tag).group(1)
        rid = re.search(r'\b\w+:id="([^"]+)"', tag).group(1)
        blaetter[unescape(name)] = (i, ziele[rid])
    return blaetter


def zeilen_ausblenden(xml, von, bis):
    """Setzt hidden="1" an den Zeilen von..bis (fehlende Zeilen werden angelegt)."""
    m = re.search(r"<sheetData\s*/>|<sheetData>(.*?)</sheetData>", xml, re.S)
    if not m:
        raise ValueError("kein <sheetData> im Blatt gefunden")
    inhalt = m.group(1) or ""
    if re.search(r"<\w+:row\b", inhalt):
        raise ValueError("Blatt verwendet XML-Präfixe – nicht unterstützt")
    zeilen = {}
    for rm in re.finditer(r"<row\b[^>]*?(?:/>|>.*?</row>)", inhalt, re.S):
        r = int(re.search(r'\br="(\d+)"', rm.group(0)).group(1))
        zeilen[r] = rm.group(0)
    for r in range(von, bis + 1):
        if r in zeilen:
            kopf = re.match(r"<row\b[^>]*?(/?>)", zeilen[r])
            start_tag = kopf.group(0)
            if re.search(r'\bhidden="[^"]*"', start_tag):
                neu = re.sub(r'\bhidden="[^"]*"', 'hidden="1"', start_tag)
            else:
                neu = start_tag[: -len(kopf.group(1))] + ' hidden="1"' + kopf.group(1)
            zeilen[r] = neu + zeilen[r][len(start_tag):]
        else:
            zeilen[r] = f'<row r="{r}" hidden="1"/>'
    neu_inhalt = "".join(zeilen[r] for r in sorted(zeilen))
    return xml[: m.start()] + f"<sheetData>{neu_inhalt}</sheetData>" + xml[m.end():]


def druckbereich_setzen(wb_xml, index, blattname, bereich):
    name = "_xlnm.Print_Area"
    ref = "'" + blattname.replace("'", "''") + "'!" + bereich
    neu = f'<definedName name="{name}" localSheetId="{index}">{escape(ref)}</definedName>'
    muster = re.compile(
        rf'<definedName\b(?=[^>]*\bname="{re.escape(name)}")(?=[^>]*\blocalSheetId="{index}")'
        r'[^>]*>.*?</definedName>', re.S)
    if muster.search(wb_xml):
        return muster.sub(lambda _: neu, wb_xml, count=1)
    leer = re.search(r"<definedNames\s*/>", wb_xml)
    if leer:
        return wb_xml[:leer.start()] + f"<definedNames>{neu}</definedNames>" + wb_xml[leer.end():]
    if "<definedNames>" in wb_xml:
        return wb_xml.replace("<definedNames>", "<definedNames>" + neu, 1)
    return wb_xml.replace("</sheets>", f"</sheets><definedNames>{neu}</definedNames>", 1)


# --------------------------------------------------------------------------
# Apple Numbers (.numbers)
# --------------------------------------------------------------------------

def numbers_zu_xlsx(eingabe, ausgabe, mit_konten):
    """Numbers-Datei: Steuerberater-Teil als neue Excel-Datei ausgeben.

    Zeilen in einer .numbers-Datei lassen sich nicht sicher ausblenden. Die
    Numbers-Datei bleibt daher unangetastet; der untere Teil wird mit den
    angezeigten Werten in eine neue Excel-Datei übernommen.
    """
    import warnings
    try:
        from numbers_parser import Document
    except ImportError:
        sys.exit("numbers-parser fehlt – bitte 'pip install numbers-parser' ausführen.")
    from openpyxl.styles import Font

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        doc = Document(eingabe)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in doc.sheets:
        for table in sheet.tables:
            def text(r, c):
                zelle = table.cell(r, c)
                try:
                    v = zelle.formatted_value
                except Exception:  # unbekanntes Zahlenformat
                    v = zelle.value
                    if isinstance(v, float):
                        v = f"{v:.2f}"
                return "" if v in (None, "None") else str(v)

            start = None
            for r in range(table.num_rows):
                t = text(r, 0)
                if MARKER_ABRECHNUNG.search(t) or (mit_konten and MARKER_KONTEN.search(t)):
                    start = r
                    break
            if start is None:
                continue
            zeilen = [[text(r, c) for c in range(table.num_cols)] for r in range(start, table.num_rows)]
            while zeilen and not any(zeilen[-1]):
                zeilen.pop()
            spalten = max((max((i for i, v in enumerate(z) if v), default=-1) for z in zeilen), default=-1) + 1

            ws = wb.create_sheet(sheet.name[:31])
            fett = Font(bold=True)
            for ri, z in enumerate(zeilen, 1):
                for ci, v in enumerate(z[:spalten], 1):
                    if v == "":
                        continue
                    roh = table.cell(start + ri - 1, ci - 1).value
                    if isinstance(roh, (int, float)) and not isinstance(roh, bool):
                        zelle = ws.cell(ri, ci, roh)
                        zelle.number_format = '#,##0.00 "€"' if "€" in v else "0.00"
                    else:
                        zelle = ws.cell(ri, ci, v)
                    if ci == 1 and (MARKER_KONTEN.search(v) or MARKER_ABRECHNUNG.search(v)
                                    or v.strip().lower() in ("weitere angaben", "position")
                                    or v.lower().startswith("summe")):
                        zelle.font = fett
            ws.column_dimensions["A"].width = 42
            for ci in range(2, spalten + 1):
                ws.column_dimensions[get_column_letter(ci)].width = 16
            ws.column_dimensions[get_column_letter(min(spalten, 5) or 1)].width = 40
            print(f"Blatt '{sheet.name}': übernommen ab Zeile {start + 1} "
                  f"('{zeilen[0][0]}'), {len(zeilen)} Zeilen")
    if not wb.worksheets:
        sys.exit("Kein Blatt mit dem Abschnitt 'Monatsabrechnung für den Steuerberater' gefunden.")
    wb.save(ausgabe)
    print(f"Gespeichert: {ausgabe}")


# --------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="Lohn ST – nur den Steuerberater-Teil der Monatsübersicht anzeigen")
    ap.add_argument("eingabe", help="Excel-Datei (.xlsx/.xlsm) oder Numbers-Datei (.numbers)")
    ap.add_argument("-o", "--ausgabe", help="Ziel-Datei (Standard: <eingabe>_Steuerberater.xlsx)")
    ap.add_argument("--blatt", action="append",
                    help="nur dieses Tabellenblatt bearbeiten (mehrfach möglich)")
    ap.add_argument("--ohne-konten", action="store_true",
                    help="auch den Block 'Urlaub / Konten' ausblenden")
    ap.add_argument("--kein-druckbereich", action="store_true",
                    help="Druckbereich nicht ändern")
    args = ap.parse_args()

    if args.eingabe.lower().endswith(".numbers"):
        ausgabe = args.ausgabe or re.sub(r"\.numbers$", "", args.eingabe, flags=re.I) + "_Steuerberater.xlsx"
        numbers_zu_xlsx(args.eingabe, ausgabe, not args.ohne_konten)
        return

    endung = ".xlsm" if args.eingabe.lower().endswith(".xlsm") else ".xlsx"
    ausgabe = args.ausgabe or re.sub(r"\.xls[xm]$", "", args.eingabe, flags=re.I) + "_Steuerberater" + endung

    wb = openpyxl.load_workbook(args.eingabe, data_only=True)
    auftraege = {}
    for ws in wb.worksheets:
        if args.blatt and ws.title not in args.blatt:
            continue
        bereich = finde_bereich(ws, not args.ohne_konten)
        if bereich:
            auftraege[ws.title] = bereich
    if not auftraege:
        sys.exit("Kein Blatt mit dem Abschnitt 'Monatsabrechnung für den Steuerberater' gefunden.")

    with zipfile.ZipFile(args.eingabe) as zin:
        pfade = blatt_pfade(zin)
        geaendert = {}
        wb_xml = zin.read("xl/workbook.xml").decode("utf-8")
        for name, (start, ende, spalte) in auftraege.items():
            index, pfad = pfade[name]
            if start > 1:
                xml = zin.read(pfad).decode("utf-8")
                geaendert[pfad] = zeilen_ausblenden(xml, 1, start - 1).encode("utf-8")
            if not args.kein_druckbereich:
                wb_xml = druckbereich_setzen(
                    wb_xml, index, name, f"$A${start}:${get_column_letter(spalte)}${ende}")
        if not args.kein_druckbereich:
            geaendert["xl/workbook.xml"] = wb_xml.encode("utf-8")

        with zipfile.ZipFile(ausgabe, "w") as zout:
            for info in zin.infolist():
                daten = geaendert.get(info.filename)
                zout.writestr(info, daten if daten is not None else zin.read(info.filename))

    for name, (start, ende, spalte) in auftraege.items():
        print(f"Blatt '{name}': Zeilen 1–{start - 1} ausgeblendet, sichtbar bleibt "
              f"A{start}:{get_column_letter(spalte)}{ende}")
    print(f"Gespeichert: {ausgabe}")


if __name__ == "__main__":
    main()
